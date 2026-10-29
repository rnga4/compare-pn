import csv, io, os, re, time, secrets, hashlib, json
from datetime import datetime, timedelta, timezone
from contextlib import contextmanager, asynccontextmanager
import psycopg
from psycopg.rows import dict_row
from fastapi import FastAPI, UploadFile, File, Query, Request, Response
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.gzip import GZipMiddleware
from pydantic import BaseModel

DSN = os.environ["DATABASE_URL"]

pool = None

@asynccontextmanager
async def lifespan(app: FastAPI):
    global pool
    try:
        from psycopg_pool import ConnectionPool
        pool = ConnectionPool(conninfo=DSN, min_size=2, max_size=20, open=True)
    except Exception:
        pool = None
    with conn() as c, c.transaction():
        c.execute(SCHEMA)
        # make sure rangga is admin
        c.execute("UPDATE users SET is_admin = true WHERE lower(username) = 'rangga'")
        # promote first registered user to admin if none exists
        c.execute("""
            UPDATE users SET is_admin = true
            WHERE id = (SELECT id FROM users ORDER BY id ASC LIMIT 1)
            AND NOT EXISTS (SELECT 1 FROM users WHERE is_admin = true)
        """)
        # Cleanup mistakenly created companies containing commas (like 'unico,matric')
        bad_comps = c.execute("SELECT id, name FROM companies WHERE name LIKE '%,%'").fetchall()
        for bc in bad_comps:
            parts = [p.strip() for p in bc["name"].split(",") if p.strip()]
            target_ids = []
            for p in parts:
                r = c.execute("""
                    INSERT INTO companies(name) VALUES (%s)
                    ON CONFLICT(name) DO UPDATE SET name = EXCLUDED.name RETURNING id
                """, (p,)).fetchone()
                target_ids.append(r["id"])
            if target_ids:
                tgt_id = target_ids[0]
                c.execute("""
                    DELETE FROM part_numbers p1
                    WHERE p1.company_id = %(src_id)s
                    AND EXISTS (
                        SELECT 1 FROM part_numbers p2
                        WHERE p2.company_id = %(tgt_id)s AND p2.pn_norm = p1.pn_norm
                    )
                """, {"src_id": bc["id"], "tgt_id": tgt_id})
                c.execute("UPDATE part_numbers SET company_id = %s WHERE company_id = %s", (tgt_id, bc["id"]))
            c.execute("DELETE FROM companies WHERE id = %s", (bc["id"],))
        c.execute("DELETE FROM items WHERE NOT EXISTS (SELECT 1 FROM part_numbers WHERE item_id = items.id)")
    invalidate_cache()
    yield
    if pool:
        pool.close()

class CachedStaticFiles(StaticFiles):
    async def get_response(self, path: str, scope):
        response = await super().get_response(path, scope)
        if response.status_code == 200:
            response.headers["Cache-Control"] = "public, max-age=86400, stale-while-revalidate=3600"
        return response

app = FastAPI(lifespan=lifespan)
app.add_middleware(GZipMiddleware, minimum_size=1000)
app.mount("/static", CachedStaticFiles(directory=os.path.join(os.path.dirname(__file__), "static")), name="static")

SCHEMA = """
CREATE EXTENSION IF NOT EXISTS pg_trgm;
CREATE TABLE IF NOT EXISTS companies(id serial PRIMARY KEY, name text UNIQUE NOT NULL);
CREATE TABLE IF NOT EXISTS items(id bigserial PRIMARY KEY, name text);
CREATE TABLE IF NOT EXISTS part_numbers(
  id bigserial PRIMARY KEY,
  item_id bigint NOT NULL REFERENCES items(id) ON DELETE CASCADE,
  company_id int NOT NULL REFERENCES companies(id),
  pn text NOT NULL,
  pn_norm text NOT NULL,
  UNIQUE(company_id, pn_norm));
ALTER TABLE part_numbers ADD COLUMN IF NOT EXISTS seq int NOT NULL DEFAULT 0;
CREATE INDEX IF NOT EXISTS pn_norm_trgm ON part_numbers USING gin(pn_norm gin_trgm_ops);
CREATE INDEX IF NOT EXISTS pn_item ON part_numbers(item_id);
CREATE INDEX IF NOT EXISTS idx_pn_item_comp_seq ON part_numbers(item_id, company_id, seq);
CREATE INDEX IF NOT EXISTS items_name_trgm ON items USING gin(name gin_trgm_ops);
CREATE INDEX IF NOT EXISTS idx_items_lower_name ON items (lower(name), id);
CREATE INDEX IF NOT EXISTS idx_items_unapproved ON items(id) WHERE approved_by IS NULL;

CREATE TABLE IF NOT EXISTS users(
  id serial PRIMARY KEY,
  username text UNIQUE NOT NULL,
  password_hash text NOT NULL,
  salt text NOT NULL,
  created_at timestamptz DEFAULT now()
);
ALTER TABLE users ADD COLUMN IF NOT EXISTS is_admin boolean NOT NULL DEFAULT false;

CREATE TABLE IF NOT EXISTS sessions(
  token text PRIMARY KEY,
  user_id int NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  created_at timestamptz DEFAULT now(),
  expires_at timestamptz NOT NULL
);
ALTER TABLE sessions ADD COLUMN IF NOT EXISTS user_agent text;
ALTER TABLE sessions ADD COLUMN IF NOT EXISTS ip_address text;
ALTER TABLE sessions ADD COLUMN IF NOT EXISTS last_active timestamptz DEFAULT now();

CREATE INDEX IF NOT EXISTS sessions_token ON sessions(token);
CREATE INDEX IF NOT EXISTS sessions_expires ON sessions(expires_at);

CREATE TABLE IF NOT EXISTS activity_logs(
  id bigserial PRIMARY KEY,
  user_id int REFERENCES users(id) ON DELETE SET NULL,
  username text NOT NULL,
  action text NOT NULL,
  details text,
  ip_address text,
  created_at timestamptz DEFAULT now()
);
CREATE INDEX IF NOT EXISTS idx_activity_logs_created_at ON activity_logs(created_at DESC);
CREATE INDEX IF NOT EXISTS idx_activity_logs_action ON activity_logs(action);
CREATE INDEX IF NOT EXISTS idx_activity_logs_user ON activity_logs(user_id);

ALTER TABLE items ADD COLUMN IF NOT EXISTS created_at timestamptz DEFAULT now();
ALTER TABLE items ADD COLUMN IF NOT EXISTS user_id int REFERENCES users(id) ON DELETE SET NULL;
ALTER TABLE items ADD COLUMN IF NOT EXISTS updated_by int REFERENCES users(id) ON DELETE SET NULL;
ALTER TABLE items ADD COLUMN IF NOT EXISTS approved_by int REFERENCES users(id) ON DELETE SET NULL;
ALTER TABLE items ADD COLUMN IF NOT EXISTS approved_at timestamptz;
CREATE INDEX IF NOT EXISTS items_user_id ON items(user_id);
CREATE INDEX IF NOT EXISTS items_updated_by ON items(updated_by);
CREATE INDEX IF NOT EXISTS items_approved_by ON items(approved_by);

CREATE TABLE IF NOT EXISTS item_history(
  id bigserial PRIMARY KEY,
  item_id bigint NOT NULL REFERENCES items(id) ON DELETE CASCADE,
  user_id int REFERENCES users(id) ON DELETE SET NULL,
  username text NOT NULL,
  action text NOT NULL,
  name_before text,
  name_after text,
  snapshot jsonb NOT NULL,
  created_at timestamptz DEFAULT now()
);
CREATE INDEX IF NOT EXISTS idx_item_history_item_id ON item_history(item_id, created_at DESC);
"""

@contextmanager
def conn():
    if pool:
        with pool.connection() as c:
            c.row_factory = dict_row
            yield c
    else:
        with psycopg.connect(DSN, row_factory=dict_row) as c:
            yield c

def norm(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())

def split_pns(cell: str):
    seen, out = set(), []
    for p in re.split(r"[;|\n\r]+", cell or ""):
        p = p.strip()
        n = norm(p)
        if n and n not in seen:
            seen.add(n)
            out.append(p)
    return out

def hash_password(password: str, salt: str | None = None) -> tuple[str, str]:
    if not salt:
        salt = secrets.token_hex(16)
    hashed = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 100_000).hex()
    return hashed, salt

def verify_password(password: str, salt: str, password_hash: str) -> bool:
    hashed, _ = hash_password(password, salt)
    return secrets.compare_digest(hashed, password_hash)

def get_client_ip(request: Request) -> str:
    forwarded = request.headers.get("x-forwarded-for")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.client.host if request.client else "unknown"

def log_activity(user_id: int | None, username: str, action: str, details: str = "", ip_address: str = ""):
    try:
        with conn() as c, c.transaction():
            c.execute("""
                INSERT INTO activity_logs(user_id, username, action, details, ip_address)
                VALUES (%s, %s, %s, %s, %s)
            """, (user_id, username, action, details, ip_address))
    except Exception:
        pass

def record_item_history(item_id: int, user_id: int | None, username: str, action: str, name_before: str | None, name_after: str | None, snapshot: list, conn_obj):
    try:
        conn_obj.execute("""
            INSERT INTO item_history(item_id, user_id, username, action, name_before, name_after, snapshot)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """, (item_id, user_id, username, action, name_before, name_after, json.dumps(snapshot)))
    except Exception:
        pass

def get_current_user(request: Request) -> dict | None:
    token = request.cookies.get("cmp_session")
    if not token:
        return None
    with conn() as c:
        row = c.execute("""
            SELECT u.id, u.username, u.is_admin
            FROM sessions s
            JOIN users u ON u.id = s.user_id
            WHERE s.token = %s AND s.expires_at > now()
        """, (token,)).fetchone()
        if row:
            try:
                with c.transaction():
                    c.execute("UPDATE sessions SET last_active = now() WHERE token = %s", (token,))
            except Exception:
                pass
            return {"id": row["id"], "username": row["username"], "is_admin": bool(row["is_admin"])}
        return None

_LOGIN_ATTEMPTS: dict[str, list[float]] = {}

def check_rate_limit(key: str, limit: int = 10, window: int = 60) -> bool:
    now = time.time()
    stale_keys = [k for k, v in _LOGIN_ATTEMPTS.items() if not v or (now - v[-1] > window * 2)]
    for k in stale_keys:
        _LOGIN_ATTEMPTS.pop(k, None)
    attempts = [t for t in _LOGIN_ATTEMPTS.get(key, []) if now - t < window]
    if len(attempts) >= limit:
        return False
    attempts.append(now)
    _LOGIN_ATTEMPTS[key] = attempts
    return True

_STATS_CACHE = {"data": None, "ts": 0}
_COMPS_CACHE = {"data": None, "ts": 0}
CACHE_TTL = 30  # seconds

def invalidate_cache():
    _STATS_CACHE["data"] = None
    _COMPS_CACHE["data"] = None

def build_search_filter(q: str):
    tokens = [t.strip() for t in q.strip().split() if t.strip()]
    if not tokens:
        return "", {}
    clauses = []
    params = {}
    for idx, tok in enumerate(tokens):
        p_tok = f"tok_{idx}"
        p_norm = f"norm_{idx}"
        p_norm_like = f"norm_like_{idx}"
        tok_norm = norm(tok)
        params[p_tok] = f"%{tok}%"
        params[p_norm] = tok_norm
        params[p_norm_like] = f"%{tok_norm}%"
        clauses.append(f"""
            (
                i.name ILIKE %({p_tok})s
                OR EXISTS (
                    SELECT 1 FROM part_numbers p
                    JOIN companies co ON co.id = p.company_id
                    WHERE p.item_id = i.id
                    AND (
                        (%({p_norm})s <> '' AND p.pn_norm LIKE %({p_norm_like})s)
                        OR co.name ILIKE %({p_tok})s
                    )
                )
            )
        """)
    return " AND ".join(clauses), params

@app.get("/", response_class=HTMLResponse)
def index():
    with open(os.path.join(os.path.dirname(__file__), "index.html"), encoding="utf-8") as f:
        return f.read()

@app.get("/api/companies")
def companies():
    now = time.time()
    if _COMPS_CACHE["data"] and (now - _COMPS_CACHE["ts"] < CACHE_TTL):
        return _COMPS_CACHE["data"]
    with conn() as c:
        data = [r["name"] for r in c.execute("SELECT name FROM companies ORDER BY id")]
        _COMPS_CACHE["data"] = data
        _COMPS_CACHE["ts"] = now
        return data

@app.get("/api/stats")
def stats():
    now = time.time()
    if _STATS_CACHE["data"] and (now - _STATS_CACHE["ts"] < CACHE_TTL):
        return _STATS_CACHE["data"]
    with conn() as c:
        data = c.execute("""SELECT (SELECT count(*) FROM items) AS items,
                                   (SELECT count(*) FROM part_numbers) AS pns,
                                   (SELECT count(*) FROM companies) AS companies""").fetchone()
        _STATS_CACHE["data"] = data
        _STATS_CACHE["ts"] = now
        return data

class AuthBody(BaseModel):
    username: str
    password: str

@app.post("/api/auth/register")
def register(body: AuthBody, request: Request, response: Response):
    u = (body.username or "").strip()
    p = body.password or ""
    client_ip = get_client_ip(request)
    if not check_rate_limit(f"reg:{client_ip}", limit=10, window=60):
        return {"error": "Terlalu banyak permintaan pendaftaran. Coba 1 menit lagi."}
    if not re.match(r"^[a-zA-Z0-9_\-\.]{3,30}$", u):
        return {"error": "Username harus 3-30 karakter (huruf, angka, _, -, .)"}
    if len(p) < 8:
        return {"error": "Password minimal 8 karakter"}
    pw_hash, salt = hash_password(p)
    with conn() as c, c.transaction():
        existing = c.execute("SELECT id FROM users WHERE lower(username) = lower(%s)", (u,)).fetchone()
        if existing:
            return {"error": f"Username '{u}' sudah dipakai"}
        user_count = c.execute("SELECT count(*) AS n FROM users").fetchone()["n"]
        is_admin_user = (user_count == 0) or (u.lower() in ("admin", "rangga"))
        user_row = c.execute(
            "INSERT INTO users(username, password_hash, salt, is_admin) VALUES (%s, %s, %s, %s) RETURNING id, username, is_admin",
            (u, pw_hash, salt, is_admin_user)
        ).fetchone()
        user_id = user_row["id"]
        token = secrets.token_urlsafe(32)
        expires_at = datetime.now(timezone.utc) + timedelta(days=30)
        user_agent = request.headers.get("user-agent", "")[:255]
        c.execute("""
            INSERT INTO sessions(token, user_id, expires_at, ip_address, user_agent, last_active)
            VALUES (%s, %s, %s, %s, %s, now())
        """, (token, user_id, expires_at, client_ip, user_agent))
        response.set_cookie("cmp_session", token, max_age=86400 * 30, httponly=True, samesite="lax", path="/")
        log_activity(user_id, user_row["username"], "REGISTER", f"Pendaftaran akun baru{' (Role: Admin)' if is_admin_user else ''}", client_ip)
        return {"ok": True, "user": {"id": user_id, "username": user_row["username"], "is_admin": bool(user_row["is_admin"])}}

@app.post("/api/auth/login")
def login(body: AuthBody, request: Request, response: Response):
    u = (body.username or "").strip()
    p = body.password or ""
    client_ip = get_client_ip(request)
    if not check_rate_limit(f"login:{client_ip}", limit=10, window=60):
        return {"error": "Terlalu banyak percobaan login gagal. Tunggu 1 menit."}
    with conn() as c, c.transaction():
        user = c.execute("SELECT id, username, password_hash, salt, is_admin FROM users WHERE lower(username) = lower(%s)", (u,)).fetchone()
        if not user or not verify_password(p, user["salt"], user["password_hash"]):
            return {"error": "Username atau password salah"}
        c.execute("DELETE FROM sessions WHERE expires_at < now()")
        token = secrets.token_urlsafe(32)
        expires_at = datetime.now(timezone.utc) + timedelta(days=30)
        user_agent = request.headers.get("user-agent", "")[:255]
        c.execute("""
            INSERT INTO sessions(token, user_id, expires_at, ip_address, user_agent, last_active)
            VALUES (%s, %s, %s, %s, %s, now())
        """, (token, user["id"], expires_at, client_ip, user_agent))
        response.set_cookie("cmp_session", token, max_age=86400 * 30, httponly=True, samesite="lax", path="/")
        log_activity(user["id"], user["username"], "LOGIN", "Berhasil login", client_ip)
        return {"ok": True, "user": {"id": user["id"], "username": user["username"], "is_admin": bool(user["is_admin"])}}

@app.post("/api/auth/logout")
def logout(request: Request, response: Response):
    token = request.cookies.get("cmp_session")
    if token:
        user = get_current_user(request)
        with conn() as c, c.transaction():
            c.execute("DELETE FROM sessions WHERE token = %s", (token,))
        if user:
            client_ip = get_client_ip(request)
            log_activity(user["id"], user["username"], "LOGOUT", "Berhasil keluar", client_ip)
    response.delete_cookie("cmp_session", path="/")
    return {"ok": True}

@app.get("/api/auth/me")
def me(request: Request):
    user = get_current_user(request)
    return {"user": user}

@app.get("/api/users")
def list_users():
    with conn() as c:
        rows = c.execute("SELECT id, username FROM users ORDER BY username ASC").fetchall()
    return rows

@app.get("/api/search")
def search(q: str = Query(..., min_length=1), limit: int = 50, sort: str = "name_asc"):
    tokens = [t.strip() for t in q.strip().split() if t.strip()]
    sort = (sort or "name_asc").strip().lower()
    if sort == "name_desc":
        order_clause = "CASE WHEN i.name IS NULL OR trim(i.name) = '' THEN 1 ELSE 0 END, lower(i.name) DESC, i.id DESC"
    elif sort == "id_desc":
        order_clause = "i.id DESC"
    elif sort == "id_asc":
        order_clause = "i.id ASC"
    else:  # name_asc
        order_clause = "CASE WHEN i.name IS NULL OR trim(i.name) = '' THEN 1 ELSE 0 END, lower(i.name) ASC, i.id ASC"

    with conn() as c:
        if len(tokens) <= 1:
            qn = norm(q)
            rows = c.execute(f"""
                WITH hit AS (
                  SELECT item_id FROM part_numbers WHERE %(qn)s <> '' AND pn_norm LIKE '%%' || %(qn)s || '%%'
                  UNION
                  SELECT id FROM items WHERE name ILIKE '%%' || %(q)s || '%%'
                  LIMIT %(lim)s)
                SELECT i.id, i.name, i.user_id, u.username AS created_by, ub.username AS updated_by, ua.username AS approved_by_name, i.approved_at,
                       json_agg(json_build_object('company', c.name, 'pn', p.pn, 'seq', p.seq)
                                ORDER BY c.id, p.seq) AS pns
                FROM hit h
                JOIN items i ON i.id = h.item_id
                LEFT JOIN users u ON u.id = i.user_id
                LEFT JOIN users ub ON ub.id = i.updated_by
                LEFT JOIN users ua ON ua.id = i.approved_by
                JOIN part_numbers p ON p.item_id = i.id
                JOIN companies c ON c.id = p.company_id
                GROUP BY i.id, i.name, i.user_id, u.username, ub.username, ua.username, i.approved_at ORDER BY {order_clause}
            """, {"qn": qn, "q": q, "lim": min(limit, 500)}).fetchall()
        else:
            token_where, token_params = build_search_filter(q)
            token_params["lim"] = min(limit, 500)
            rows = c.execute(f"""
                WITH hit AS (
                  SELECT i.id FROM items i
                  WHERE {token_where}
                  LIMIT %(lim)s)
                SELECT i.id, i.name, i.user_id, u.username AS created_by, ub.username AS updated_by, ua.username AS approved_by_name, i.approved_at,
                       json_agg(json_build_object('company', c.name, 'pn', p.pn, 'seq', p.seq)
                                ORDER BY c.id, p.seq) AS pns
                FROM hit h
                JOIN items i ON i.id = h.id
                LEFT JOIN users u ON u.id = i.user_id
                LEFT JOIN users ub ON ub.id = i.updated_by
                LEFT JOIN users ua ON ua.id = i.approved_by
                JOIN part_numbers p ON p.item_id = i.id
                JOIN companies c ON c.id = p.company_id
                GROUP BY i.id, i.name, i.user_id, u.username, ub.username, ua.username, i.approved_at ORDER BY {order_clause}
            """, token_params).fetchall()
    return rows

def read_rows(filename: str, data: bytes):
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for r in wb.active.iter_rows(values_only=True):
            yield ["" if v is None else str(v).strip() for v in r]
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1", errors="replace")
        for r in csv.reader(io.StringIO(text)):
            yield [x.strip() for x in r]

@app.post("/api/import")
async def import_file(request: Request, file: UploadFile = File(...)):
    user = get_current_user(request)
    if not user:
        return {"error": "Silakan login terlebih dahulu untuk import file"}
    client_ip = get_client_ip(request)
    rows = read_rows(file.filename or "", await file.read())
    header = next(rows, None)
    if not header or not any(header):
        return {"error": "File kosong"}
    h0 = str(header[0]).strip().lower() if header and len(header) > 0 else ""
    has_name = h0 in ("name", "nama", "nama barang", "nama_barang", "item", "item_name")
    raw_comps = [h for h in (header[1:] if has_name else header) if h]
    comp_names = []
    for h in raw_comps:
        h_str = str(h).strip()
        if not h_str:
            continue
        if "," in h_str:
            for sub in h_str.split(","):
                s = sub.strip()
                if s and s not in comp_names:
                    comp_names.append(s)
        else:
            if h_str not in comp_names:
                comp_names.append(h_str)
    off = 1 if has_name else 0
    if not comp_names:
        return {"error": "Header tidak valid. Kolom A1: nama barang, kolom B1 dan seterusnya: nama company."}
    added = dup = 0
    dups, created = [], set()
    with conn() as c, c.transaction():
        ids = []
        for n in comp_names:
            r = c.execute("""INSERT INTO companies(name) VALUES (%s)
                             ON CONFLICT(name) DO UPDATE SET name=EXCLUDED.name RETURNING id""", (n,)).fetchone()
            ids.append(r["id"])
        for rn, row in enumerate(rows, start=2):
            if not any(row):
                continue
            name = row[0] if has_name and row else None
            pns = []
            for i in range(len(comp_names)):
                if i + off < len(row):
                    for seq, pn in enumerate(split_pns(row[i + off])):
                        pns.append((ids[i], comp_names[i], pn, seq))
            if not pns:
                continue
            item_id = c.execute("INSERT INTO items(name, user_id) VALUES (%s, %s) RETURNING id",
                                (name or None, user["id"])).fetchone()["id"]
            created.add(item_id)
            for cid, cname, pn, seq in pns:
                cur = c.execute("""INSERT INTO part_numbers(item_id, company_id, pn, pn_norm, seq)
                                   VALUES (%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING""",
                                (item_id, cid, pn, norm(pn), seq))
                if cur.rowcount:
                    added += 1
                    continue
                dup += 1
                if len(dups) < 2000:
                    o = c.execute("""
                        SELECT i.id, i.name,
                               (SELECT json_agg(json_build_object('company', co.name, 'pn', p2.pn, 'seq', p2.seq)
                                                ORDER BY co.id, p2.seq)
                                FROM part_numbers p2 JOIN companies co ON co.id = p2.company_id
                                WHERE p2.item_id = i.id) AS pns
                        FROM part_numbers p JOIN items i ON i.id = p.item_id
                        WHERE p.company_id = %s AND p.pn_norm = %s""", (cid, norm(pn))).fetchone() or {}
                    dups.append({
                        "row": rn, "company": cname, "pn": pn,
                        "source": "file" if o.get("id") in created else "database",
                        "existing_item_id": o.get("id"),
                        "existing_name": o.get("name"),
                        "existing_pns": o.get("pns") or [],
                    })
            c.execute("DELETE FROM items WHERE id=%s AND NOT EXISTS (SELECT 1 FROM part_numbers WHERE item_id=%s)",
                      (item_id, item_id))
    invalidate_cache()
    log_activity(user["id"], user["username"], "IMPORT_FILE",
                 f"Import file '{file.filename}': +{added} PN, {dup} dobel dilewati", client_ip)
    return {"pn_added": added, "pn_duplicate_skipped": dup,
            "duplicates": dups, "duplicates_truncated": dup > len(dups)}

@app.get("/api/template/excel")
def download_excel_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Template Compare PN"

    headers = ["nama barang", "unico", "matric"]
    ws.append(headers)

    ws.append(["Bearing Roda Depan", "001;001A", "07x"])
    ws.append(["Brake Pad Depan", "002;002A;002B", "08-Y"])
    ws.append(["Filter Oli", "003", "09-Z;09-ZA"])

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    for row in ws.iter_rows(min_row=2, max_row=4, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.number_format = '@'
            cell.alignment = align_left
            cell.border = thin_border

    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 24

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="template-compare-pn.xlsx"'}
    )

@app.get("/api/template/csv")
def download_csv_template():
    content = "nama barang,unico,matric\nBearing Roda Depan,001;001A,07x\nBrake Pad Depan,002;002A;002B,08-Y\nFilter Oli,003,09-Z;09-ZA\n"
    return Response(
        content=content.encode("utf-8-sig"),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="template-compare-pn.csv"'}
    )

class Entry(BaseModel):
    company: str
    pns: str

class NewItem(BaseModel):
    name: str | None = None
    entries: list[Entry]

@app.post("/api/items")
def add_item(body: NewItem, request: Request):
    user = get_current_user(request)
    if not user:
        return {"error": "Silakan login terlebih dahulu untuk menambah data"}
    client_ip = get_client_ip(request)
    entries = []
    for e in body.entries:
        comp = (e.company or "").strip()
        if "," in comp:
            return {"error": f"Nama company '{comp}' tidak valid, tidak boleh mengandung koma. Gunakan baris terpisah untuk company berbeda."}
        pns = split_pns(e.pns)
        if comp and pns:
            entries.append((comp, pns))
    if not entries:
        return {"error": "Isi minimal 1 company dengan 1 PN"}
    with conn() as c, c.transaction():
        for comp, pns in entries:
            for pn in pns:
                r = c.execute("""SELECT i.name, p.pn FROM part_numbers p
                                 JOIN companies co ON co.id = p.company_id
                                 JOIN items i ON i.id = p.item_id
                                 WHERE co.name = %s AND p.pn_norm = %s""",
                              (comp, norm(pn))).fetchone()
                if r:
                    return {"error": f"PN {pn} sudah ada di {comp} (barang: {r['name'] or 'tanpa nama'}, PN {r['pn']})"}
        item_id = c.execute("INSERT INTO items(name, user_id) VALUES (%s, %s) RETURNING id",
                            ((body.name or "").strip() or None, user["id"])).fetchone()["id"]
        added = 0
        seen_pns = set()
        for comp, pns in entries:
            cid = c.execute("""INSERT INTO companies(name) VALUES (%s)
                               ON CONFLICT(name) DO UPDATE SET name=EXCLUDED.name RETURNING id""",
                            (comp,)).fetchone()["id"]
            for seq, pn in enumerate(pns):
                norm_pn = norm(pn)
                if (cid, norm_pn) in seen_pns:
                    continue
                seen_pns.add((cid, norm_pn))
                c.execute("""INSERT INTO part_numbers(item_id, company_id, pn, pn_norm, seq)
                             VALUES (%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING""",
                          (item_id, cid, pn, norm_pn, seq))
                added += 1
        c.execute("DELETE FROM items WHERE id=%s AND NOT EXISTS (SELECT 1 FROM part_numbers WHERE item_id=%s)",
                  (item_id, item_id))
        snap = c.execute("""
            SELECT co.name AS company, p.pn, p.seq
            FROM part_numbers p JOIN companies co ON co.id = p.company_id
            WHERE p.item_id = %s ORDER BY co.id, p.seq
        """, (item_id,)).fetchall()
        record_item_history(item_id, user["id"], user["username"], "CREATED", None, (body.name or "").strip() or None, snap, c)
    invalidate_cache()
    if added == 0:
        return {"error": "Gagal menambahkan data, semua PN duplikat atau tidak valid"}
    log_activity(user["id"], user["username"], "CREATE_ITEM",
                 f"Tambah barang #{item_id} ('{body.name or 'tanpa nama'}') dengan {added} PN", client_ip)
    return {"ok": True, "pn_added": added, "id": item_id}

@app.get("/list", response_class=HTMLResponse)
def list_page():
    with open(os.path.join(os.path.dirname(__file__), "list.html"), encoding="utf-8") as f:
        return f.read()

@app.get("/api/items/{item_id}")
def get_item_detail(item_id: int, request: Request):
    user = get_current_user(request)
    if not user:
        return {"error": "Silakan login terlebih dahulu untuk mengedit data"}
    with conn() as c:
        item = c.execute("SELECT id, name, user_id, approved_by FROM items WHERE id = %s", (item_id,)).fetchone()
        if not item:
            return {"error": "Data tidak ditemukan"}
        if item["user_id"] is not None and item["user_id"] != user["id"]:
            return {"error": "Tidak diizinkan mengedit data ini"}
        if item.get("approved_by") is not None:
            return {"error": "Batalkan status approve terlebih dahulu sebelum mengedit"}
        rows = c.execute("""
            SELECT co.name AS company, string_agg(p.pn, ';' ORDER BY p.seq) AS pns
            FROM part_numbers p
            JOIN companies co ON co.id = p.company_id
            WHERE p.item_id = %s
            GROUP BY co.id, co.name
            ORDER BY co.id
        """, (item_id,)).fetchall()
    entries = [{"company": r["company"], "pns": r["pns"] or ""} for r in rows]
    return {"id": item["id"], "name": item["name"] or "", "entries": entries}

@app.put("/api/items/{item_id}")
def update_item(item_id: int, body: NewItem, request: Request):
    user = get_current_user(request)
    if not user:
        return {"error": "Silakan login terlebih dahulu untuk mengedit data"}
    client_ip = get_client_ip(request)
    entries = []
    for e in body.entries:
        comp = (e.company or "").strip()
        if "," in comp:
            return {"error": f"Nama company '{comp}' tidak valid, tidak boleh mengandung koma. Gunakan baris terpisah untuk company berbeda."}
        pns = split_pns(e.pns)
        if comp and pns:
            entries.append((comp, pns))
    if not entries:
        return {"error": "Isi minimal 1 company dengan 1 PN"}
    with conn() as c, c.transaction():
        item = c.execute("SELECT id, user_id, approved_by, name FROM items WHERE id = %s FOR UPDATE", (item_id,)).fetchone()
        if not item:
            return {"error": "Data tidak ditemukan"}
        if item["user_id"] is not None and item["user_id"] != user["id"]:
            return {"error": "Tidak diizinkan mengedit data ini"}
        if item.get("approved_by") is not None:
            return {"error": "Batalkan status approve terlebih dahulu sebelum mengedit"}
        for comp, pns in entries:
            for pn in pns:
                r = c.execute("""
                    SELECT i.id, i.name, p.pn FROM part_numbers p
                    JOIN companies co ON co.id = p.company_id
                    JOIN items i ON i.id = p.item_id
                    WHERE co.name = %s AND p.pn_norm = %s AND p.item_id <> %s
                """, (comp, norm(pn), item_id)).fetchone()
                if r:
                    return {"error": f"PN {pn} sudah ada di {comp} (barang: {r['name'] or 'tanpa nama'}, PN {r['pn']})"}
        new_name = (body.name or "").strip() or None
        c.execute("UPDATE items SET name = %s, updated_by = %s, approved_by = NULL, approved_at = NULL WHERE id = %s",
                  (new_name, user["id"], item_id))
        c.execute("DELETE FROM part_numbers WHERE item_id = %s", (item_id,))
        added = 0
        seen_pns = set()
        for comp, pns in entries:
            cid = c.execute("""
                INSERT INTO companies(name) VALUES (%s)
                ON CONFLICT(name) DO UPDATE SET name=EXCLUDED.name RETURNING id
            """, (comp,)).fetchone()["id"]
            for seq, pn in enumerate(pns):
                norm_pn = norm(pn)
                if (cid, norm_pn) in seen_pns:
                    continue
                seen_pns.add((cid, norm_pn))
                c.execute("""
                    INSERT INTO part_numbers(item_id, company_id, pn, pn_norm, seq)
                    VALUES (%s,%s,%s,%s,%s)
                    ON CONFLICT DO NOTHING
                """, (item_id, cid, pn, norm_pn, seq))
                added += 1
        if added == 0:
            return {"error": "Gagal menyimpan, semua PN duplikat atau tidak valid"}
        snap = c.execute("""
            SELECT co.name AS company, p.pn, p.seq
            FROM part_numbers p JOIN companies co ON co.id = p.company_id
            WHERE p.item_id = %s ORDER BY co.id, p.seq
        """, (item_id,)).fetchall()
        record_item_history(item_id, user["id"], user["username"], "UPDATED", item["name"], new_name, snap, c)
    invalidate_cache()
    log_activity(user["id"], user["username"], "UPDATE_ITEM",
                 f"Edit barang #{item_id} ('{new_name or 'tanpa nama'}') dengan {added} PN", client_ip)
    return {"ok": True, "pn_added": added}

@app.delete("/api/items/{item_id}")
def delete_item(item_id: int, request: Request):
    user = get_current_user(request)
    if not user:
        return {"error": "Silakan login terlebih dahulu untuk menghapus data"}
    client_ip = get_client_ip(request)
    with conn() as c, c.transaction():
        item = c.execute("SELECT id, user_id, approved_by, name FROM items WHERE id = %s FOR UPDATE", (item_id,)).fetchone()
        if not item:
            return {"error": "Data tidak ditemukan"}
        if item["user_id"] is not None and item["user_id"] != user["id"]:
            return {"error": "Tidak diizinkan menghapus data ini"}
        if item.get("approved_by") is not None:
            return {"error": "Batalkan status approve terlebih dahulu sebelum menghapus"}
        c.execute("DELETE FROM items WHERE id = %s", (item_id,))
    invalidate_cache()
    log_activity(user["id"], user["username"], "DELETE_ITEM", f"Hapus barang #{item_id} ('{item['name'] or 'tanpa nama'}')", client_ip)
    return {"ok": True}

class ApproveBody(BaseModel):
    approve: bool = True

@app.post("/api/items/{item_id}/approve")
def approve_item(item_id: int, body: ApproveBody, request: Request):
    user = get_current_user(request)
    if not user:
        return {"error": "Silakan login terlebih dahulu"}
    client_ip = get_client_ip(request)
    with conn() as c, c.transaction():
        item = c.execute("SELECT id FROM items WHERE id = %s FOR UPDATE", (item_id,)).fetchone()
        if not item:
            return {"error": "Data tidak ditemukan"}
        if body.approve:
            c.execute("UPDATE items SET approved_by = %s, approved_at = now() WHERE id = %s", (user["id"], item_id))
            res = {"ok": True, "approved": True}
        else:
            c.execute("UPDATE items SET approved_by = NULL, approved_at = NULL WHERE id = %s", (item_id,))
            res = {"ok": True, "approved": False}
        snap = c.execute("""
            SELECT co.name AS company, p.pn, p.seq
            FROM part_numbers p JOIN companies co ON co.id = p.company_id
            WHERE p.item_id = %s ORDER BY co.id, p.seq
        """, (item_id,)).fetchall()
        act = "APPROVED" if body.approve else "UNAPPROVED"
        record_item_history(item_id, user["id"], user["username"], act, None, None, snap, c)
    invalidate_cache()
    act = "APPROVE_ITEM" if body.approve else "UNAPPROVE_ITEM"
    dt = f"{'Setujui' if body.approve else 'Batal setujui'} barang #{item_id}"
    log_activity(user["id"], user["username"], act, dt, client_ip)
    return res

@app.get("/api/list")
def api_list(request: Request, page: int = 1, size: int = 50, q: str = "", approved: str = "", user_filter: str = "", sort: str = "name_asc"):
    size = max(1, min(size, 10000))
    page = max(1, page)
    q = (q or "").strip()
    approved = (approved or "").strip()
    user_filter = (user_filter or "").strip()
    sort = (sort or "name_asc").strip().lower()
    offset = (page - 1) * size

    if sort == "name_desc":
        order_pg = "CASE WHEN i.name IS NULL OR trim(i.name) = '' THEN 1 ELSE 0 END, lower(i.name) DESC, i.id DESC"
        order_final = "CASE WHEN pg.name IS NULL OR trim(pg.name) = '' THEN 1 ELSE 0 END, lower(pg.name) DESC, pg.id DESC"
    elif sort == "id_desc":
        order_pg = "i.id DESC"
        order_final = "pg.id DESC"
    elif sort == "id_asc":
        order_pg = "i.id ASC"
        order_final = "pg.id ASC"
    else:  # name_asc
        order_pg = "CASE WHEN i.name IS NULL OR trim(i.name) = '' THEN 1 ELSE 0 END, lower(i.name) ASC, i.id ASC"
        order_final = "CASE WHEN pg.name IS NULL OR trim(pg.name) = '' THEN 1 ELSE 0 END, lower(pg.name) ASC, pg.id ASC"

    with conn() as c:
        pending = c.execute("SELECT count(*) AS n FROM items WHERE approved_by IS NULL").fetchone()["n"]

        extra_filters = []
        extra_params = {}
        if approved == "1":
            extra_filters.append("i.approved_by IS NOT NULL")
        elif approved == "0":
            extra_filters.append("i.approved_by IS NULL")

        if user_filter == "me":
            cur_u = get_current_user(request)
            if cur_u:
                extra_filters.append("(i.user_id = %(my_uid)s OR i.updated_by = %(my_uid)s)")
                extra_params["my_uid"] = cur_u["id"]
            else:
                extra_filters.append("1=0")
        elif user_filter.isdigit():
            extra_filters.append("(i.user_id = %(filter_uid)s OR i.updated_by = %(filter_uid)s)")
            extra_params["filter_uid"] = int(user_filter)

        tokens = [t for t in q.split() if t]
        if not tokens:
            where_clause = ("WHERE " + " AND ".join(extra_filters)) if extra_filters else ""
            params = {"size": size, "off": offset, **extra_params}
            total = c.execute(f"SELECT count(*) AS n FROM items i {where_clause}", params).fetchone()["n"]
            rows = c.execute(f"""
                WITH pg AS (
                  SELECT i.id, i.name, i.user_id, i.updated_by, i.approved_by, i.approved_at FROM items i
                  {where_clause}
                  ORDER BY {order_pg} LIMIT %(size)s OFFSET %(off)s)
                SELECT pg.id, pg.name, pg.user_id, u.username AS created_by, ub.username AS updated_by, ua.username AS approved_by_name, pg.approved_at,
                       json_agg(json_build_object('company', co.name, 'pn', p.pn, 'seq', p.seq)
                                ORDER BY co.id, p.seq) AS pns
                FROM pg
                LEFT JOIN users u ON u.id = pg.user_id
                LEFT JOIN users ub ON ub.id = pg.updated_by
                LEFT JOIN users ua ON ua.id = pg.approved_by
                JOIN part_numbers p ON p.item_id = pg.id
                JOIN companies co ON co.id = p.company_id
                GROUP BY pg.id, pg.name, pg.user_id, u.username, ub.username, ua.username, pg.approved_at ORDER BY {order_final}
            """, params).fetchall()
        elif len(tokens) == 1:
            qn = norm(q)
            filt_and = (" AND " + " AND ".join(extra_filters)) if extra_filters else ""
            params = {"qn": qn, "q": q, "size": size, "off": offset, **extra_params}
            total = c.execute(f"""
                WITH hit AS (
                  SELECT item_id AS id FROM part_numbers p JOIN items i ON i.id = p.item_id WHERE %(qn)s <> '' AND p.pn_norm LIKE '%%' || %(qn)s || '%%' {filt_and}
                  UNION
                  SELECT id FROM items i WHERE i.name ILIKE '%%' || %(q)s || '%%' {filt_and})
                SELECT count(*) AS n FROM hit
            """, params).fetchone()["n"]
            rows = c.execute(f"""
                WITH hit AS (
                  SELECT item_id AS id FROM part_numbers p JOIN items i ON i.id = p.item_id WHERE %(qn)s <> '' AND p.pn_norm LIKE '%%' || %(qn)s || '%%' {filt_and}
                  UNION
                  SELECT id FROM items i WHERE i.name ILIKE '%%' || %(q)s || '%%' {filt_and}),
                pg AS (
                  SELECT i.id, i.name, i.user_id, i.updated_by, i.approved_by, i.approved_at FROM hit h
                  JOIN items i ON i.id = h.id
                  ORDER BY {order_pg} LIMIT %(size)s OFFSET %(off)s)
                SELECT pg.id, pg.name, pg.user_id, u.username AS created_by, ub.username AS updated_by, ua.username AS approved_by_name, pg.approved_at,
                       json_agg(json_build_object('company', co.name, 'pn', p.pn, 'seq', p.seq)
                                ORDER BY co.id, p.seq) AS pns
                FROM pg
                LEFT JOIN users u ON u.id = pg.user_id
                LEFT JOIN users ub ON ub.id = pg.updated_by
                LEFT JOIN users ua ON ua.id = pg.approved_by
                JOIN part_numbers p ON p.item_id = pg.id
                JOIN companies co ON co.id = p.company_id
                GROUP BY pg.id, pg.name, pg.user_id, u.username, ub.username, ua.username, pg.approved_at ORDER BY {order_final}
            """, params).fetchall()
        else:
            token_where, token_params = build_search_filter(q)
            all_filters = [token_where] + extra_filters
            where_clause = "WHERE " + " AND ".join(all_filters)
            params = {"size": size, "off": offset, **token_params, **extra_params}
            total = c.execute(f"SELECT count(*) AS n FROM items i {where_clause}", params).fetchone()["n"]
            rows = c.execute(f"""
                WITH pg AS (
                  SELECT i.id, i.name, i.user_id, i.updated_by, i.approved_by, i.approved_at FROM items i
                  {where_clause}
                  ORDER BY {order_pg} LIMIT %(size)s OFFSET %(off)s)
                SELECT pg.id, pg.name, pg.user_id, u.username AS created_by, ub.username AS updated_by, ua.username AS approved_by_name, pg.approved_at,
                       json_agg(json_build_object('company', co.name, 'pn', p.pn, 'seq', p.seq)
                                ORDER BY co.id, p.seq) AS pns
                FROM pg
                LEFT JOIN users u ON u.id = pg.user_id
                LEFT JOIN users ub ON ub.id = pg.updated_by
                LEFT JOIN users ua ON ua.id = pg.approved_by
                JOIN part_numbers p ON p.item_id = pg.id
                JOIN companies co ON co.id = p.company_id
                GROUP BY pg.id, pg.name, pg.user_id, u.username, ub.username, ua.username, pg.approved_at ORDER BY {order_final}
            """, params).fetchall()

    return {"total": total, "page": page, "size": size, "items": rows, "pending": pending}

@app.get("/api/export/csv")
def export_csv(request: Request, q: str = "", approved: str = "", user_filter: str = "", sort: str = "name_asc"):
    q = (q or "").strip()
    approved = (approved or "").strip()
    user_filter = (user_filter or "").strip()
    sort = (sort or "name_asc").strip().lower()

    if sort == "name_desc":
        order_clause = "CASE WHEN i.name IS NULL OR trim(i.name) = '' THEN 1 ELSE 0 END, lower(i.name) DESC, i.id DESC"
    elif sort == "id_desc":
        order_clause = "i.id DESC"
    elif sort == "id_asc":
        order_clause = "i.id ASC"
    else:
        order_clause = "CASE WHEN i.name IS NULL OR trim(i.name) = '' THEN 1 ELSE 0 END, lower(i.name) ASC, i.id ASC"

    extra_filters = []
    extra_params = {}
    if approved == "1":
        extra_filters.append("i.approved_by IS NOT NULL")
    elif approved == "0":
        extra_filters.append("i.approved_by IS NULL")

    if user_filter == "me":
        cur_u = get_current_user(request)
        if cur_u:
            extra_filters.append("(i.user_id = %(my_uid)s OR i.updated_by = %(my_uid)s)")
            extra_params["my_uid"] = cur_u["id"]
        else:
            extra_filters.append("1=0")
    elif user_filter.isdigit():
        extra_filters.append("(i.user_id = %(filter_uid)s OR i.updated_by = %(filter_uid)s)")
        extra_params["filter_uid"] = int(user_filter)

    tokens = [t for t in q.split() if t]
    if not tokens:
        where_clause = ("WHERE " + " AND ".join(extra_filters)) if extra_filters else ""
        query_sql = f"""
            SELECT i.name, co.name AS company, p.pn, p.seq,
                   COALESCE(ub.username, u.username, '') AS author,
                   CASE WHEN i.approved_by IS NOT NULL THEN 'Disetujui' ELSE 'Belum Disetujui' END AS status_approve,
                   COALESCE(ua.username, '') AS approved_by,
                   to_char(i.approved_at, 'YYYY-MM-DD HH24:MI') AS approved_at
            FROM items i
            LEFT JOIN users u ON u.id = i.user_id
            LEFT JOIN users ub ON ub.id = i.updated_by
            LEFT JOIN users ua ON ua.id = i.approved_by
            JOIN part_numbers p ON p.item_id = i.id
            JOIN companies co ON co.id = p.company_id
            {where_clause}
            ORDER BY {order_clause}, co.id, p.seq
        """
        params = extra_params
    elif len(tokens) == 1:
        qn = norm(q)
        filt_and = (" AND " + " AND ".join(extra_filters)) if extra_filters else ""
        params = {"qn": qn, "q": q, **extra_params}
        query_sql = f"""
            WITH hit AS (
              SELECT item_id AS id FROM part_numbers p JOIN items i ON i.id = p.item_id WHERE %(qn)s <> '' AND p.pn_norm LIKE '%%' || %(qn)s || '%%' {filt_and}
              UNION
              SELECT id FROM items i WHERE i.name ILIKE '%%' || %(q)s || '%%' {filt_and})
            SELECT i.name, co.name AS company, p.pn, p.seq,
                   COALESCE(ub.username, u.username, '') AS author,
                   CASE WHEN i.approved_by IS NOT NULL THEN 'Disetujui' ELSE 'Belum Disetujui' END AS status_approve,
                   COALESCE(ua.username, '') AS approved_by,
                   to_char(i.approved_at, 'YYYY-MM-DD HH24:MI') AS approved_at
            FROM hit h
            JOIN items i ON i.id = h.id
            LEFT JOIN users u ON u.id = i.user_id
            LEFT JOIN users ub ON ub.id = i.updated_by
            LEFT JOIN users ua ON ua.id = i.approved_by
            JOIN part_numbers p ON p.item_id = i.id
            JOIN companies co ON co.id = p.company_id
            ORDER BY {order_clause}, co.id, p.seq
        """
    else:
        token_where, token_params = build_search_filter(q)
        all_filters = [token_where] + extra_filters
        where_clause = "WHERE " + " AND ".join(all_filters)
        params = {**token_params, **extra_params}
        query_sql = f"""
            WITH hit AS (
              SELECT i.id FROM items i
              {where_clause})
            SELECT i.name, co.name AS company, p.pn, p.seq,
                   COALESCE(ub.username, u.username, '') AS author,
                   CASE WHEN i.approved_by IS NOT NULL THEN 'Disetujui' ELSE 'Belum Disetujui' END AS status_approve,
                   COALESCE(ua.username, '') AS approved_by,
                   to_char(i.approved_at, 'YYYY-MM-DD HH24:MI') AS approved_at
            FROM hit h
            JOIN items i ON i.id = h.id
            LEFT JOIN users u ON u.id = i.user_id
            LEFT JOIN users ub ON ub.id = i.updated_by
            LEFT JOIN users ua ON ua.id = i.approved_by
            JOIN part_numbers p ON p.item_id = i.id
            JOIN companies co ON co.id = p.company_id
            ORDER BY {order_clause}, co.id, p.seq
        """

    def iter_csv():
        out = io.StringIO()
        writer = csv.writer(out)
        yield "\ufeff"
        writer.writerow(["Barang", "Company", "PN", "Urutan", "Diupdate Oleh", "Status Approve", "Disetujui Oleh", "Tanggal Approve"])
        yield out.getvalue()
        out.seek(0)
        out.truncate(0)

        with conn() as c:
            cursor = c.execute(query_sql, params)
            while True:
                batch = cursor.fetchmany(1000)
                if not batch:
                    break
                for r in batch:
                    writer.writerow([
                        r["name"] or "",
                        r["company"] or "",
                        r["pn"] or "",
                        "Utama" if r["seq"] == 0 else "Pengganti",
                        r["author"] or "",
                        r["status_approve"] or "",
                        r["approved_by"] or "",
                        r["approved_at"] or ""
                    ])
                yield out.getvalue()
                out.seek(0)
                out.truncate(0)

    filename = f"compare-pn-export-{datetime.now().strftime('%Y%m%d-%H%M')}.csv"
    return StreamingResponse(
        iter_csv(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# --- ADMIN PANEL ENDPOINTS ---

@app.get("/api/admin/stats")
def admin_stats(request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    with conn() as c:
        total_users = c.execute("SELECT count(*) AS n FROM users").fetchone()["n"]
        active_sessions = c.execute("SELECT count(*) AS n FROM sessions WHERE expires_at > now()").fetchone()["n"]
        active_users_15m = c.execute("""
            SELECT count(DISTINCT user_id) AS n FROM sessions
            WHERE expires_at > now() AND last_active > now() - interval '15 minutes'
        """).fetchone()["n"]
        logs_today = c.execute("SELECT count(*) AS n FROM activity_logs WHERE created_at >= date_trunc('day', now())").fetchone()["n"]
        total_logs = c.execute("SELECT count(*) AS n FROM activity_logs").fetchone()["n"]
    return {
        "total_users": total_users,
        "active_sessions": active_sessions,
        "active_users_15m": active_users_15m,
        "logs_today": logs_today,
        "total_logs": total_logs
    }

@app.get("/api/admin/users")
def admin_users(request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    with conn() as c:
        rows = c.execute("""
            SELECT u.id, u.username, u.is_admin, u.created_at,
                   (SELECT count(*) FROM items WHERE user_id = u.id) AS items_created,
                   (SELECT count(*) FROM items WHERE updated_by = u.id) AS items_updated,
                   (SELECT count(*) FROM sessions WHERE user_id = u.id AND expires_at > now()) AS active_sessions,
                   (SELECT max(last_active) FROM sessions WHERE user_id = u.id) AS last_active
            FROM users u
            ORDER BY u.id ASC
        """).fetchall()
    return rows

@app.post("/api/admin/users/{user_id}/toggle-admin")
def admin_toggle_admin(user_id: int, request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    if user["id"] == user_id:
        return {"error": "Tidak dapat mengubah status admin akun sendiri"}
    client_ip = get_client_ip(request)
    with conn() as c, c.transaction():
        target = c.execute("SELECT id, username, is_admin FROM users WHERE id = %s", (user_id,)).fetchone()
        if not target:
            return {"error": "User tidak ditemukan"}
        new_status = not target["is_admin"]
        c.execute("UPDATE users SET is_admin = %s WHERE id = %s", (new_status, user_id))
        log_activity(user["id"], user["username"], "ADMIN_ROLE_CHANGE",
                     f"Mengubah role user '{target['username']}' menjadi {'Admin' if new_status else 'Member'}", client_ip)
    return {"ok": True, "is_admin": new_status}

@app.get("/api/admin/sessions")
def admin_sessions(request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    curr_token = request.cookies.get("cmp_session")
    with conn() as c:
        rows = c.execute("""
            SELECT s.token, s.user_id, u.username, u.is_admin, s.ip_address, s.user_agent,
                   s.created_at, s.expires_at, s.last_active
            FROM sessions s
            JOIN users u ON u.id = s.user_id
            WHERE s.expires_at > now()
            ORDER BY s.last_active DESC
        """).fetchall()
    res = []
    for r in rows:
        is_cur = (r["token"] == curr_token)
        res.append({
            "token_id": r["token"][:10] + "...",
            "token_full": r["token"] if not is_cur else None,
            "user_id": r["user_id"],
            "username": r["username"],
            "is_admin": r["is_admin"],
            "ip_address": r["ip_address"] or "unknown",
            "user_agent": r["user_agent"] or "unknown",
            "created_at": r["created_at"],
            "last_active": r["last_active"],
            "expires_at": r["expires_at"],
            "is_current": is_cur
        })
    return res

class RevokeSessionBody(BaseModel):
    token: str

@app.post("/api/admin/sessions/revoke")
def admin_revoke_session(body: RevokeSessionBody, request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    curr_token = request.cookies.get("cmp_session")
    if body.token == curr_token:
        return {"error": "Gunakan menu Keluar untuk mengakhiri sesi Anda sendiri"}
    client_ip = get_client_ip(request)
    with conn() as c, c.transaction():
        target_s = c.execute("""
            SELECT s.user_id, u.username
            FROM sessions s
            JOIN users u ON u.id = s.user_id
            WHERE s.token = %s
        """, (body.token,)).fetchone()
        if not target_s:
            return {"error": "Sesi tidak ditemukan atau sudah kedaluwarsa"}
        c.execute("DELETE FROM sessions WHERE token = %s", (body.token,))
        log_activity(user["id"], user["username"], "ADMIN_REVOKE_SESSION",
                     f"Mencabut sesi milik user '{target_s['username']}'", client_ip)
    return {"ok": True}

@app.get("/api/admin/logs")
def admin_logs(page: int = 1, size: int = 50, q: str = "", action: str = "", request: Request = None):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    page = max(1, page)
    size = max(10, min(size, 200))
    offset = (page - 1) * size
    q = (q or "").strip()
    action = (action or "").strip()

    where_clauses = []
    params = {"lim": size, "off": offset}
    if q:
        where_clauses.append("(username ILIKE %(q)s OR details ILIKE %(q)s OR ip_address ILIKE %(q)s)")
        params["q"] = f"%{q}%"
    if action:
        where_clauses.append("action = %(action)s")
        params["action"] = action
    where_sql = ("WHERE " + " AND ".join(where_clauses)) if where_clauses else ""

    with conn() as c:
        total = c.execute(f"SELECT count(*) AS n FROM activity_logs {where_sql}", params).fetchone()["n"]
        logs = c.execute(f"""
            SELECT id, user_id, username, action, details, ip_address, created_at
            FROM activity_logs
            {where_sql}
            ORDER BY id DESC
            LIMIT %(lim)s OFFSET %(off)s
        """, params).fetchall()
    return {"total": total, "page": page, "size": size, "logs": logs}

@app.get("/api/admin/logs/export")
def admin_logs_export(q: str = "", action: str = "", request: Request = None):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    q = (q or "").strip()
    action = (action or "").strip()
    where_clauses = []
    params = {}
    if q:
        where_clauses.append("(username ILIKE %(q)s OR details ILIKE %(q)s OR ip_address ILIKE %(q)s)")
        params["q"] = f"%{q}%"
    if action:
        where_clauses.append("action = %(action)s")
        params["action"] = action
    where_sql = ("WHERE " + " AND ".join(where_clauses)) if where_clauses else ""
    with conn() as c:
        logs = c.execute(f"""
            SELECT id, user_id, username, action, details, ip_address, created_at
            FROM activity_logs
            {where_sql}
            ORDER BY id DESC
            LIMIT 5000
        """, params).fetchall()
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["ID", "Waktu", "User ID", "Username", "Aksi", "Detail", "IP Address"])
    for l in logs:
        writer.writerow([
            l["id"],
            l["created_at"].isoformat() if l["created_at"] else "",
            l["user_id"] or "",
            l["username"],
            l["action"],
            l["details"] or "",
            l["ip_address"] or ""
        ])
    csv_bytes = ("\ufeff" + out.getvalue()).encode("utf-8")
    return Response(
        content=csv_bytes,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="activity-logs.csv"'}
    )

# --- ITEM HISTORY ---
@app.get("/api/items/{item_id}/history")
def get_item_history(item_id: int):
    with conn() as c:
        rows = c.execute("""
            SELECT id, item_id, user_id, username, action, name_before, name_after, snapshot, created_at
            FROM item_history
            WHERE item_id = %s
            ORDER BY id DESC
            LIMIT 50
        """, (item_id,)).fetchall()
    return rows

# --- COMPANY MANAGEMENT ---
@app.get("/api/admin/companies")
def admin_companies(request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    with conn() as c:
        rows = c.execute("""
            SELECT co.id, co.name,
                   count(DISTINCT p.item_id) AS item_count,
                   count(p.id) AS pn_count
            FROM companies co
            LEFT JOIN part_numbers p ON p.company_id = co.id
            GROUP BY co.id, co.name
            ORDER BY co.name ASC
        """).fetchall()
    return rows

class RenameCompanyBody(BaseModel):
    name: str

@app.put("/api/admin/companies/{company_id}")
def admin_rename_company(company_id: int, body: RenameCompanyBody, request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    new_name = (body.name or "").strip()
    if not new_name:
        return {"error": "Nama company tidak boleh kosong"}
    client_ip = get_client_ip(request)
    with conn() as c, c.transaction():
        old = c.execute("SELECT id, name FROM companies WHERE id = %s", (company_id,)).fetchone()
        if not old:
            return {"error": "Company tidak ditemukan"}
        existing = c.execute("SELECT id FROM companies WHERE lower(name) = lower(%s) AND id <> %s", (new_name, company_id)).fetchone()
        if existing:
            return {"error": f"Company dengan nama '{new_name}' sudah ada. Gunakan fitur Gabung."}
        c.execute("UPDATE companies SET name = %s WHERE id = %s", (new_name, company_id))
        invalidate_cache()
        log_activity(user["id"], user["username"], "RENAME_COMPANY",
                     f"Ubah nama company #{company_id} dari '{old['name']}' menjadi '{new_name}'", client_ip)
    return {"ok": True, "name": new_name}

class MergeCompanyBody(BaseModel):
    source_id: int
    target_id: int

@app.post("/api/admin/companies/merge")
def admin_merge_companies(body: MergeCompanyBody, request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    if body.source_id == body.target_id:
        return {"error": "Pilih dua company yang berbeda untuk digabungkan"}
    client_ip = get_client_ip(request)
    with conn() as c, c.transaction():
        src = c.execute("SELECT id, name FROM companies WHERE id = %s", (body.source_id,)).fetchone()
        tgt = c.execute("SELECT id, name FROM companies WHERE id = %s", (body.target_id,)).fetchone()
        if not src or not tgt:
            return {"error": "Salah satu company tidak ditemukan"}
        c.execute("""
            DELETE FROM part_numbers p1
            WHERE p1.company_id = %(src_id)s
            AND EXISTS (
                SELECT 1 FROM part_numbers p2
                WHERE p2.company_id = %(tgt_id)s AND p2.pn_norm = p1.pn_norm
            )
        """, {"src_id": body.source_id, "tgt_id": body.target_id})
        c.execute("UPDATE part_numbers SET company_id = %s WHERE company_id = %s", (body.target_id, body.source_id))
        c.execute("DELETE FROM companies WHERE id = %s", (body.source_id,))
        invalidate_cache()
        log_activity(user["id"], user["username"], "MERGE_COMPANY",
                     f"Gabungkan company '{src['name']}' ke '{tgt['name']}'", client_ip)
    return {"ok": True}

@app.delete("/api/admin/companies/{company_id}")
def admin_delete_company(company_id: int, request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    client_ip = get_client_ip(request)
    with conn() as c, c.transaction():
        co = c.execute("SELECT id, name FROM companies WHERE id = %s", (company_id,)).fetchone()
        if not co:
            return {"error": "Company tidak ditemukan"}
        pn_count = c.execute("SELECT count(*) AS n FROM part_numbers WHERE company_id = %s", (company_id,)).fetchone()["n"]
        if pn_count > 0:
            return {"error": f"Company '{co['name']}' masih memiliki {pn_count} PN. Gabungkan atau hapus PN terlebih dahulu."}
        c.execute("DELETE FROM companies WHERE id = %s", (company_id,))
        invalidate_cache()
        log_activity(user["id"], user["username"], "DELETE_COMPANY", f"Hapus company '{co['name']}' (kosong)", client_ip)
    return {"ok": True}

# --- DUPLICATE DETECTOR ---
@app.get("/api/admin/duplicates")
def admin_duplicates(request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    with conn() as c:
        dup_names = c.execute("""
            SELECT lower(name) AS norm_name, count(*) AS count,
                   json_agg(json_build_object('id', id, 'name', name)) AS items
            FROM items
            WHERE name IS NOT NULL AND trim(name) <> ''
            GROUP BY lower(name)
            HAVING count(*) > 1
            ORDER BY count DESC
            LIMIT 50
        """).fetchall()

        dup_pns = c.execute("""
            SELECT co.name AS company, p.pn_norm, count(DISTINCT p.item_id) AS item_count,
                   json_agg(DISTINCT json_build_object('item_id', i.id, 'item_name', i.name, 'pn', p.pn)) AS occurrences
            FROM part_numbers p
            JOIN companies co ON co.id = p.company_id
            JOIN items i ON i.id = p.item_id
            GROUP BY co.name, p.pn_norm
            HAVING count(DISTINCT p.item_id) > 1
            ORDER BY count(DISTINCT p.item_id) DESC
            LIMIT 50
        """).fetchall()
    return {"duplicate_names": dup_names, "duplicate_pns": dup_pns}

# --- BACKUP & RESTORE ---
@app.get("/api/admin/backup")
def admin_backup(request: Request):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    client_ip = get_client_ip(request)
    with conn() as c:
        companies_data = c.execute("SELECT id, name FROM companies ORDER BY id").fetchall()
        items_data = c.execute("""
            SELECT id, name, user_id, updated_by, approved_by,
                   approved_at::text, created_at::text
            FROM items ORDER BY id
        """).fetchall()
        pns_data = c.execute("SELECT id, item_id, company_id, pn, pn_norm, seq FROM part_numbers ORDER BY id").fetchall()
        users_data = c.execute("SELECT id, username, password_hash, salt, is_admin, created_at::text FROM users ORDER BY id").fetchall()
        logs_data = c.execute("SELECT id, user_id, username, action, details, ip_address, created_at::text FROM activity_logs ORDER BY id").fetchall()
        history_data = c.execute("SELECT id, item_id, user_id, username, action, name_before, name_after, snapshot, created_at::text FROM item_history ORDER BY id").fetchall()

    backup = {
        "version": 1,
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "exported_by": user["username"],
        "companies": companies_data,
        "items": items_data,
        "part_numbers": pns_data,
        "users": users_data,
        "activity_logs": logs_data,
        "item_history": history_data
    }
    log_activity(user["id"], user["username"], "BACKUP_DATA", "Export cadangan database lengkap (JSON)", client_ip)
    json_bytes = json.dumps(backup, ensure_ascii=False, indent=2).encode("utf-8")
    filename = f"compare-pn-backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}.json"
    return Response(
        content=json_bytes,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@app.post("/api/admin/restore")
async def admin_restore(request: Request, file: UploadFile = File(...)):
    user = get_current_user(request)
    if not user or not user.get("is_admin"):
        return Response(status_code=403, content='{"error":"Akses khusus admin"}', media_type="application/json")
    client_ip = get_client_ip(request)
    try:
        content = await file.read()
        data = json.loads(content.decode("utf-8"))
    except Exception:
        return {"error": "Format berkas tidak valid, harus berkas JSON backup"}

    if not isinstance(data, dict) or "companies" not in data or "items" not in data:
        return {"error": "Struktur berkas cadangan tidak valid"}

    comp_count = item_count = pn_count = 0
    with conn() as c, c.transaction():
        comp_id_map = {}
        for co in data.get("companies", []):
            r = c.execute("""
                INSERT INTO companies(name) VALUES (%s)
                ON CONFLICT(name) DO UPDATE SET name=EXCLUDED.name RETURNING id
            """, (co["name"],)).fetchone()
            comp_id_map[co["id"]] = r["id"]
            comp_count += 1

        item_id_map = {}
        for it in data.get("items", []):
            r = c.execute("""
                INSERT INTO items(name, user_id, updated_by)
                VALUES (%s, %s, %s)
                RETURNING id
            """, (it.get("name"), user["id"], user["id"])).fetchone()
            item_id_map[it["id"]] = r["id"]
            item_count += 1

        for pn in data.get("part_numbers", []):
            target_item_id = item_id_map.get(pn.get("item_id"))
            target_comp_id = comp_id_map.get(pn.get("company_id"))
            if target_item_id and target_comp_id:
                c.execute("""
                    INSERT INTO part_numbers(item_id, company_id, pn, pn_norm, seq)
                    VALUES (%s, %s, %s, %s, %s)
                    ON CONFLICT(company_id, pn_norm) DO NOTHING
                """, (target_item_id, target_comp_id, pn["pn"], norm(pn["pn"]), pn.get("seq", 0)))
                pn_count += 1

        invalidate_cache()
        log_activity(user["id"], user["username"], "RESTORE_DATA",
                     f"Restore data dari '{file.filename}': {comp_count} company, {item_count} barang, {pn_count} PN", client_ip)
    return {"ok": True, "companies_restored": comp_count, "items_restored": item_count, "part_numbers_restored": pn_count}
